"""Single text-extraction interface. Swappable backends.

Scanned PDFs and images go through an optional ``vision`` hook — a callable
``(bytes, mime) -> str`` such as :meth:`charta.core.llm.LLMClient.describe_image`.
Without it, extraction degrades to a filename-only stub.
"""
from __future__ import annotations

import mimetypes
import pathlib
from dataclasses import dataclass
from typing import Callable, Sequence

from ..sensitive import redact

VISION_MIME = "image/png"
VisionFn = Callable[[bytes, str], str]
IMAGE_SUFFIXES = (".png", ".jpg", ".jpeg", ".bmp", ".webp", ".gif", ".tif", ".tiff")


@dataclass
class ExtractResult:
    text: str
    pages: int = 1
    method: str = "none"


def extract_text(path: str | pathlib.Path, vision: VisionFn | None = None) -> ExtractResult:
    p = pathlib.Path(path)
    suf = p.suffix.lower()
    try:
        if suf in (".txt", ".md", ".csv", ".json", ".toml", ".yaml", ".yml"):
            return ExtractResult(p.read_text(encoding="utf-8", errors="replace")[:20000], 1, "text")
        if suf == ".pdf":
            return _extract_pdf(p, vision)
        if suf == ".docx":
            return _extract_docx(p)
        if suf in (".xlsx", ".xlsm", ".xltx"):
            return _extract_xlsx(p)
        if suf in IMAGE_SUFFIXES:
            return _extract_image(p, vision)
        return ExtractResult(f"[unsupported: {p.name}]", 1, "stub")
    except FileNotFoundError:
        return ExtractResult("", 0, "missing")
    except Exception as e:  # noqa: BLE001
        return ExtractResult(f"[extract-error: {e}]", 1, "error")


def _read_image(p: pathlib.Path) -> bytes:
    return p.read_bytes()


def _extract_image(p: pathlib.Path, vision: VisionFn | None) -> ExtractResult:
    if vision is not None:
        mime = mimetypes.guess_type(p.name)[0] or VISION_MIME
        text = _safe_vision(vision, _read_image(p), mime)
        if text.strip():
            return ExtractResult(text[:20000], 1, "vision")
    return ExtractResult(f"[image file: {p.name}]", 1, "image-stub")


def _safe_vision(vision: VisionFn, data: bytes, mime: str) -> str:
    """Vision calls are best-effort: any failure must not break extraction."""
    try:
        return redact(vision(data, mime) or "")
    except Exception:  # noqa: BLE001
        return ""


def _extract_pdf(p: pathlib.Path, vision: VisionFn | None = None) -> ExtractResult:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return ExtractResult(f"[pdf without PyMuPDF installed: {p.name}]", 1, "missing-dep")
    try:
        doc = fitz.open(p)
        chunks = [page.get_text() for page in doc]
        text = "\n".join(chunks)[:30000]
        pages = len(doc)
        if text.strip():
            doc.close()
            return ExtractResult(text, pages, "pymupdf")
        if vision is None:
            doc.close()
            return ExtractResult(
                f"[scanned pdf, {pages} page(s): {p.name} — needs vision model]", pages, "vision-needed")
        rendered = "\n".join(_safe_vision(vision, _render_page(page), VISION_MIME)
                             for page in doc)[:30000]
        doc.close()
        if rendered.strip():
            return ExtractResult(rendered, pages, "vision")
        return ExtractResult(f"[scanned pdf, {pages} page(s), vision empty: {p.name}]", pages, "vision-empty")
    except Exception as e:  # noqa: BLE001
        return ExtractResult(f"[pdf-error: {e}]", 1, "error")


def _render_page(page) -> bytes:  # page: fitz.Page
    return page.get_pixmap(dpi=150).tobytes("png")


def extract_many(paths: Sequence[str | pathlib.Path], vision: VisionFn | None = None) -> list[ExtractResult]:
    """Extract a batch, reusing the same vision hook (keeps call sites tidy)."""
    return [extract_text(p, vision) for p in paths]


def _extract_docx(p: pathlib.Path) -> ExtractResult:
    # Prefer python-docx when it is installed (preserves paragraphs).
    try:
        import docx

        doc = docx.Document(str(p))
        return ExtractResult("\n".join(par.text for par in doc.paragraphs)[:20000], 1, "docx")
    except ImportError:
        pass  # fall through to the stdlib-only reader below
    except Exception as e:  # noqa: BLE001
        return ExtractResult(f"[docx-error: {e}]", 1, "error")
    # Fallback: a .docx is just a zip of XML — read the body text directly.
    try:
        import re
        import zipfile

        with zipfile.ZipFile(p) as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="replace")
        text = re.sub(r"<[^>]+>", " ", xml)
        text = " ".join(text.split())
        if text.strip():
            return ExtractResult(text[:20000], 1, "docx-xml")
    except Exception as e:  # noqa: BLE001
        return ExtractResult(f"[docx-error: {e}]", 1, "error")
    return ExtractResult(f"[docx without python-docx: {p.name}]", 1, "missing-dep")


def _extract_xlsx(p: pathlib.Path) -> ExtractResult:
    try:
        import openpyxl
    except ImportError:
        return ExtractResult(f"[xlsx without openpyxl: {p.name}]", 1, "missing-dep")
    try:
        wb = openpyxl.load_workbook(str(p), data_only=True)
        lines: list[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" | ".join(cells))
        text = "\n".join(lines)
        if text.strip():
            return ExtractResult(text[:20000], len(wb.worksheets) or 1, "xlsx")
    except Exception as e:  # noqa: BLE001
        return ExtractResult(f"[xlsx-error: {e}]", 1, "error")
    return ExtractResult(f"[empty xlsx: {p.name}]", 1, "xlsx-empty")
